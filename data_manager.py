import os
import json
import re
import openpyxl
from typing import Dict, List, Any, Optional

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
SURVEY_STORE_PATH = os.path.join(DATA_DIR, "survey_store.json")
COMMENTS_STORE_PATH = os.path.join(DATA_DIR, "comments_store.json")
CURRICULUM_STORE_PATH = os.path.join(DATA_DIR, "curriculum_store.json")

os.makedirs(DATA_DIR, exist_ok=True)


from datetime import datetime, timedelta


def excel_serial_to_date_string(serial) -> str:
    """엑셀 시리얼 번호(예: 46104)를 YYYY-MM-DD 문자열로 변환"""
    try:
        val = float(serial)
        if 20000 <= val <= 65000:
            d = datetime(1899, 12, 30) + timedelta(days=val)
            return d.strftime("%Y-%m-%d")
    except (ValueError, TypeError):
        pass
    return ""


def parse_schedule(schedule_raw: str):
    """
    [YYYY-MM-DD ~ YYYY-MM-DD] N차 또는 제N차 파싱
    엑셀 시리얼 번호(46104, 46230, 46286 등) 자동 변환 지원
    """
    if not schedule_raw:
        return {"raw": "", "start_date": "", "end_date": "", "month": 0, "round_name": "", "display_period": ""}
    
    s = str(schedule_raw).strip()

    # 1. 엑셀 시리얼 번호 단독인 경우 (예: "46104" 또는 "46104.0")
    serial_date = excel_serial_to_date_string(s)
    if serial_date:
        month = int(serial_date.split('-')[1])
        return {
            "raw": s,
            "start_date": serial_date,
            "end_date": serial_date,
            "month": month,
            "round_name": "",
            "display_period": serial_date
        }

    # 2. [시리얼 ~ 시리얼] N차 형태
    m_serial = re.search(r'\[\s*(\d{5}(?:\.\d+)?)\s*~\s*(\d{5}(?:\.\d+)?)\s*\]\s*(.*)', s)
    if m_serial:
        s1, s2, round_name = m_serial.groups()
        d1 = excel_serial_to_date_string(s1) or s1
        d2 = excel_serial_to_date_string(s2) or s2
        round_name = round_name.strip()
        month = int(d1.split('-')[1]) if '-' in d1 else 0
        return {
            "raw": s,
            "start_date": d1,
            "end_date": d2,
            "month": month,
            "round_name": round_name,
            "display_period": f"{d1} ~ {d2}"
        }

    # 3. [YYYY-MM-DD ~ YYYY-MM-DD] N차 또는 제N차 파싱
    m = re.search(r'\[\s*(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})\s*\]\s*(.*)', s)
    if m:
        start_date, end_date, round_name = m.groups()
        round_name = round_name.strip()
        month = int(start_date.split('-')[1])
        return {
            "raw": s,
            "start_date": start_date,
            "end_date": end_date,
            "month": month,
            "round_name": round_name,
            "display_period": f"{start_date} ~ {end_date}"
        }
    
    # 4. 일자만 있는 경우
    m2 = re.search(r'(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})', s)
    if m2:
        start_date, end_date = m2.groups()
        month = int(start_date.split('-')[1])
        return {
            "raw": s,
            "start_date": start_date,
            "end_date": end_date,
            "month": month,
            "round_name": "",
            "display_period": f"{start_date} ~ {end_date}"
        }

    # 5. YYYY-MM-DD 단독인 경우
    m3 = re.search(r'(\d{4}-\d{2}-\d{2})', s)
    if m3:
        d = m3.group(1)
        month = int(d.split('-')[1])
        return {
            "raw": s,
            "start_date": d,
            "end_date": d,
            "month": month,
            "round_name": "",
            "display_period": d
        }
    
    return {"raw": s, "start_date": "", "end_date": "", "month": 0, "round_name": "", "display_period": s}


def safe_float(v, default=0.0):
    if v is None:
        return default
    try:
        return round(float(v), 2)
    except (ValueError, TypeError):
        return default


def safe_int(v, default=0):
    if v is None:
        return default
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return default


class DataManager:
    def __init__(self):
        self.records: List[Dict[str, Any]] = []
        self.comments: Dict[str, Any] = {}
        self.curriculums: Dict[str, str] = {}
        self.load_all()

    def load_all(self):
        """저장된 JSON 또는 기본 로우데이터 파일에서 데이터 로드"""
        if os.path.exists(SURVEY_STORE_PATH):
            try:
                with open(SURVEY_STORE_PATH, "r", encoding="utf-8") as f:
                    self.records = json.load(f)
            except Exception as e:
                print(f"Error loading survey store: {e}")
                self.records = []

        # 저장소가 비어있으면 기본 엑셀에서 로드
        if not self.records:
            default_raw_file = os.path.join(BASE_DIR, "과정별설문결과_2026-09-23_로우데이터.xlsx")
            if os.path.exists(default_raw_file):
                self.import_raw_excel(default_raw_file)

        # 주관식 의견 로드
        if os.path.exists(COMMENTS_STORE_PATH):
            try:
                with open(COMMENTS_STORE_PATH, "r", encoding="utf-8") as f:
                    self.comments = json.load(f)
            except Exception as e:
                print(f"Error loading comments: {e}")
                self.comments = {}

        # 커리큘럼 로드
        if os.path.exists(CURRICULUM_STORE_PATH):
            try:
                with open(CURRICULUM_STORE_PATH, "r", encoding="utf-8") as f:
                    self.curriculums = json.load(f)
            except Exception as e:
                print(f"Error loading curriculums: {e}")
                self.curriculums = {}

    def save_survey_store(self):
        with open(SURVEY_STORE_PATH, "w", encoding="utf-8") as f:
            json.dump(self.records, f, ensure_ascii=False, indent=2)

    def save_comments_store(self):
        with open(COMMENTS_STORE_PATH, "w", encoding="utf-8") as f:
            json.dump(self.comments, f, ensure_ascii=False, indent=2)

    def save_curriculum_store(self):
        with open(CURRICULUM_STORE_PATH, "w", encoding="utf-8") as f:
            json.dump(self.curriculums, f, ensure_ascii=False, indent=2)

    def import_raw_excel(self, file_path: str, replace: bool = False) -> int:
        """로우데이터 엑셀 파일을 읽어와 레코드 병합 및 저장 (지능형 복합 헤더 감지, 제외어 필터, 자가 치유)"""
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        # 1. 헤더 행 위치 지능형 탐색 (상위 15행 스캔하여 매칭 점수 계산)
        header_keywords = ['과정명', '과정', '과정정보', '교육일정', '일정', '강사명', '강사', '발송수', '응답수', '과정만족도', '추천지수', '만족도평균', '교재완성도']
        best_header_row_idx = 2
        max_score = -1

        for r in range(1, min(15, ws.max_row + 1)):
            row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, min(25, ws.max_column + 1))]
            score = 0
            for val in row_vals:
                if not val:
                    continue
                for kw in header_keywords:
                    if val == kw:
                        score += 3
                    elif kw in val:
                        score += 1
            if score > max_score:
                max_score = score
                best_header_row_idx = r

        if max_score < 2:
            best_header_row_idx = 2 if ws.max_row > 2 else 1

        # 2단 복합 헤더 구조 구축 (상위 행 + 하위 행 결합)
        num_cols = ws.max_column
        col_headers = []

        for c in range(1, num_cols + 1):
            top = str(ws.cell(best_header_row_idx - 1, c).value or '').strip() if best_header_row_idx > 1 else ''
            bot = str(ws.cell(best_header_row_idx, c).value or '').strip()

            # 상위 병합 셀 좌측 값 전파
            if not top and best_header_row_idx > 1:
                for left_c in range(c - 1, 0, -1):
                    prev_top = str(ws.cell(best_header_row_idx - 1, left_c).value or '').strip()
                    if prev_top and any(k in prev_top for k in ['과정정보', '설문대상', '강사만족도', '불편사항', '선호 교육형태', '수강동기']):
                        top = prev_top
                        break

            combined = f"{top}_{bot}" if top and bot else (bot or top)
            col_headers.append({"col": c, "top": top, "bot": bot, "combined": combined})

        def find_col(exact=None, contains=None, exclude=None, default_c=1):
            exact = exact or []
            contains = contains or []
            exclude = exclude or []

            # 1단계: 완전 일치
            for h in col_headers:
                for kw in exact:
                    if h["bot"] == kw or h["top"] == kw or h["combined"] == kw:
                        return h["col"]

            # 2단계: 제외어 검사 후 포함 일치
            for kw in contains:
                for h in col_headers:
                    text = h["combined"]
                    if not text:
                        continue
                    if any(ex in text for ex in exclude):
                        continue
                    if kw in text:
                        return h["col"]

            return default_c

        c_course = find_col(
            exact=['과정명', '교육과정명', '과정'],
            contains=['과정명', '교육과정', '과정'],
            exclude=['코드', '번호', '정보', '유형', '구분', '시간', '만족도', '차수', '비용', '강사'],
            default_c=1
        )
        c_sched = find_col(
            exact=['교육일정(차수)', '교육일정', '일정(차수)', '교육기간', '연수기간'],
            contains=['교육일정', '교육기간', '일정', '연수기간'],
            exclude=['코드', '번호', '구분', '불편', '시간', '항목', '불편사항'],
            default_c=2
        )
        c_region = find_col(
            exact=['지역', '교육장소', '장소'],
            contains=['지역', '장소', '캠퍼스'],
            exclude=['코드', '번호'],
            default_c=3
        )
        c_inst = find_col(
            exact=['강사명', '교수명', '강사'],
            contains=['강사명', '교수명', '강사'],
            exclude=['코드', '번호', '만족도', '평균', '전문성', '전달', '강의', '료', '확정', '평가'],
            default_c=4
        )
        c_hours = find_col(
            exact=['강의시간', '교육시간', '시간'],
            contains=['강의시간', '교육시간', '시간'],
            exclude=['시작', '종료', '코드'],
            default_c=5
        )
        c_sent = find_col(
            exact=['발송수', '발송건수', '설문발송수'],
            contains=['발송'],
            exclude=['일정', '일자'],
            default_c=6
        )
        c_resp = find_col(
            exact=['응답수', '응답건수', '설문응답수', '설문인원', '참여인원'],
            contains=['응답', '설문인원', '참여인원'],
            exclude=['율', '비율'],
            default_c=7
        )
        c_nps = find_col(
            exact=['추천지수(NPS)', '추천지수', 'NPS', '순추천고객지수'],
            contains=['추천지수', 'NPS', '순추천'],
            exclude=['사유', '이유', '의견', '추천인', '동료추천', '담당자추천'],
            default_c=8
        )
        c_course_sat = find_col(
            exact=['과정만족도', '교육과정만족도', '과정내용만족도', '과정만족'],
            contains=['과정만족', '과정내용만족', '교육과정만족'],
            exclude=['강사', '교재', '시설', '환경', '불편'],
            default_c=9
        )
        c_inst_avg = find_col(
            exact=['만족도평균', '강사만족도평균', '강사만족도'],
            contains=['만족도평균', '강사만족'],
            exclude=['과정'],
            default_c=10
        )
        c_expertise = find_col(
            exact=['강의전문성', '강의내용', '전문성'],
            contains=['전문성', '강의내용'],
            exclude=['코드'],
            default_c=11
        )
        c_delivery = find_col(
            exact=['전달능력', '전달력', '강의전달'],
            contains=['전달'],
            exclude=['코드'],
            default_c=12
        )
        c_practical = find_col(
            exact=['교육효과성', '실무활용도', '활용도', '실무적용도'],
            contains=['실무활용', '교육효과', '활용도', '효과성'],
            exclude=['코드'],
            default_c=13
        )
        c_textbook = find_col(
            exact=['교재완성도', '교재만족도', '교재품질', '교재'],
            contains=['교재완성', '교재품질', '교재'],
            exclude=['코드', '번호', '과정코드'],
            default_c=14
        )

        if replace:
            self.records = []

        imported_count = 0
        existing_keys = {
            f"{r.get('course_name')}_{r.get('start_date') or r.get('schedule_raw')}_{r.get('instructor_name')}": idx 
            for idx, r in enumerate(self.records)
        }

        for row_idx in range(best_header_row_idx + 1, ws.max_row + 1):
            course_name = ws.cell(row_idx, c_course).value
            if not course_name or str(course_name).strip() == '' or str(course_name).strip() in ['과정명', '과정정보']:
                continue
            
            course_name = str(course_name).strip()
            schedule_raw = str(ws.cell(row_idx, c_sched).value or "").strip()
            region = str(ws.cell(row_idx, c_region).value or "").strip()
            instructor_name = str(ws.cell(row_idx, c_inst).value or "").strip()

            # [자가치유 1] 강사명과 교육일정이 뒤바뀐 경우 감지 및 교정
            def is_date_like(val_str: str) -> bool:
                if not val_str:
                    return False
                if re.search(r'^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}', val_str) or re.search(r'^\[\s*\d{4}', val_str):
                    return True
                if excel_serial_to_date_string(val_str):
                    return True
                return False

            if is_date_like(instructor_name) and not is_date_like(schedule_raw):
                schedule_raw, instructor_name = instructor_name, schedule_raw
            elif is_date_like(instructor_name):
                # 강사명에 일정이 들어갔는데 schedule_raw도 일정이면 인근 컬럼에서 한글 강사명 탐색
                found_inst = ""
                for c in range(1, min(10, ws.max_column + 1)):
                    if c not in (c_course, c_sched):
                        val = str(ws.cell(row_idx, c).value or "").strip()
                        if re.match(r'^[가-힣]{2,4}$', val) and val not in ['서울', '부산', '대구', '대전', '광주', '인천', '울산', '경기', '온라인', '비대면']:
                            found_inst = val
                            break
                instructor_name = found_inst
            
            sched_info = parse_schedule(schedule_raw)

            # [자가치유 2] 발송수, 응답수, 만족도, NPS 자동 교정
            raw_sent = safe_int(ws.cell(row_idx, c_sent).value)
            raw_resp = safe_int(ws.cell(row_idx, c_resp).value)
            raw_nps = safe_float(ws.cell(row_idx, c_nps).value)
            raw_sat = safe_float(ws.cell(row_idx, c_course_sat).value)

            # 만족도가 5.0 초과인 경우 (예: 발송수 7이 만족도로 들어간 현상)
            if raw_sat > 5.0:
                real_sat = 0.0
                for c in range(7, 14):
                    v = safe_float(ws.cell(row_idx, c).value)
                    if 1.0 <= v <= 5.0:
                        real_sat = v
                        break
                if raw_sent == 0 and raw_sat > 5.0:
                    raw_sent = int(raw_sat)
                raw_sat = real_sat

            # NPS가 응답수(5, 6 등 소액 정수)와 일치하고 실제 NPS는 다른 열에 있는 경우
            if raw_nps <= 10.0 and raw_resp > 0 and round(raw_nps) == raw_resp:
                for c in range(6, 12):
                    v = safe_float(ws.cell(row_idx, c).value)
                    if (v > 10.0 or v == 0.0) and v != raw_sent and v != raw_resp:
                        raw_nps = v
                        break

            # [자가치유 3] 교재완성도가 5.0 초과(예: 과정코드 306,796)인 경우 차단 및 인접 정상 만족도 복원
            raw_textbook = safe_float(ws.cell(row_idx, c_textbook).value)
            if raw_textbook > 5.0 or raw_textbook < 0:
                real_tb = 0.0
                for c in range(9, 15):
                    v = safe_float(ws.cell(row_idx, c).value)
                    if 1.0 <= v <= 5.0 and c != c_course_sat:
                        real_tb = v
                raw_textbook = real_tb

            def sanitize_score(v_in):
                val = safe_float(v_in)
                return 0.0 if (val > 5.0 or val < 0) else val

            raw_inst_avg = sanitize_score(ws.cell(row_idx, c_inst_avg).value)
            raw_expertise = sanitize_score(ws.cell(row_idx, c_expertise).value)
            raw_delivery = sanitize_score(ws.cell(row_idx, c_delivery).value)
            raw_practical = sanitize_score(ws.cell(row_idx, c_practical).value)

            record = {
                "course_name": course_name,
                "schedule_raw": schedule_raw,
                "start_date": sched_info["start_date"],
                "end_date": sched_info["end_date"],
                "month": sched_info["month"],
                "round_name": sched_info["round_name"],
                "display_period": sched_info["display_period"],
                "region": region,
                "instructor_name": instructor_name,
                "hours": safe_int(ws.cell(row_idx, c_hours).value),
                "sent_count": raw_sent,
                "respondent_count": raw_resp,
                "nps": raw_nps,
                "course_satisfaction": raw_sat,
                "instructor_satisfaction_avg": raw_inst_avg,
                "teaching_expertise": raw_expertise,
                "delivery_skill": raw_delivery,
                "practical_use": raw_practical,
                "textbook_quality": raw_textbook,
                "complaints": {
                    "정보": safe_int(ws.cell(row_idx, 15).value),
                    "절차": safe_int(ws.cell(row_idx, 16).value),
                    "운영": safe_int(ws.cell(row_idx, 17).value),
                    "환경": safe_int(ws.cell(row_idx, 18).value),
                    "일정": safe_int(ws.cell(row_idx, 19).value),
                    "기타": safe_int(ws.cell(row_idx, 20).value),
                },
                "motives": {
                    "SNS": safe_int(ws.cell(row_idx, 21).value),
                    "이메일": safe_int(ws.cell(row_idx, 22).value),
                    "홈페이지": safe_int(ws.cell(row_idx, 23).value),
                    "인쇄물": safe_int(ws.cell(row_idx, 24).value),
                    "인터넷": safe_int(ws.cell(row_idx, 25).value),
                    "담당자추천": safe_int(ws.cell(row_idx, 26).value),
                    "동료추천": safe_int(ws.cell(row_idx, 27).value),
                    "기타": safe_int(ws.cell(row_idx, 28).value),
                },
                "positions": {
                    "사원": safe_int(ws.cell(row_idx, 29).value),
                    "대리": safe_int(ws.cell(row_idx, 30).value),
                    "과장": safe_int(ws.cell(row_idx, 31).value),
                    "차장": safe_int(ws.cell(row_idx, 32).value),
                    "부팀장": safe_int(ws.cell(row_idx, 33).value),
                    "임원": safe_int(ws.cell(row_idx, 34).value),
                },
                "preferred_formats": {
                    "오프라인": safe_int(ws.cell(row_idx, 35).value),
                    "비대면": safe_int(ws.cell(row_idx, 36).value),
                    "이러닝": safe_int(ws.cell(row_idx, 37).value),
                    "플립러닝": safe_int(ws.cell(row_idx, 38).value),
                },
                "course_code": str(ws.cell(row_idx, 39).value or ""),
                "round_code": str(ws.cell(row_idx, 40).value or ""),
                "instructor_code": str(ws.cell(row_idx, 41).value or ""),
            }

            key = f"{course_name}_{sched_info['start_date'] or schedule_raw}_{instructor_name}"
            if key in existing_keys:
                self.records[existing_keys[key]] = record
            else:
                existing_keys[key] = len(self.records)
                self.records.append(record)
            
            imported_count += 1

        self.save_survey_store()
        return imported_count

    def get_filter_options(self) -> Dict[str, Any]:
        """검색용 전체 과정 목록, 강사 목록 반환"""
        courses = sorted(list({r["course_name"] for r in self.records}))
        instructors = sorted(list({r["instructor_name"] for r in self.records if r["instructor_name"]}))
        
        # 과정별 강사 및 일정 매핑
        course_details = {}
        for r in self.records:
            c = r["course_name"]
            if c not in course_details:
                course_details[c] = {"instructors": set(), "schedules": []}
            if r["instructor_name"]:
                course_details[c]["instructors"].add(r["instructor_name"])
            course_details[c]["schedules"].append({
                "raw": r["schedule_raw"],
                "start_date": r["start_date"],
                "month": r["month"],
                "round_name": r["round_name"],
                "display": r["display_period"]
            })
            
        for c in course_details:
            course_details[c]["instructors"] = sorted(list(course_details[c]["instructors"]))
            # 날짜순 정렬
            course_details[c]["schedules"].sort(key=lambda x: x["start_date"])

        return {
            "courses": courses,
            "instructors": instructors,
            "course_details": course_details,
            "total_records": len(self.records)
        }

    def get_report_data(self, course_name: str, instructor_name: Optional[str] = None, mode: str = "instructor") -> Dict[str, Any]:
        """
        특정 과정 및 (선택적) 강사에 대한 보고서 데이터 생성
        mode: 'instructor' (강사별만족도) 또는 'course' (교육운영결과보고서)
        """
        # 1. 일치하는 레코드 필터링
        filtered = [r for r in self.records if r["course_name"] == course_name]
        if instructor_name and mode == "instructor":
            filtered = [r for r in filtered if r["instructor_name"] == instructor_name]

        if not filtered:
            return {"error": "해당 조건의 데이터가 존재하지 않습니다."}

        # 2. 일정 시작일 순 정렬
        filtered.sort(key=lambda x: (x["start_date"], x["schedule_raw"]))

        # 3. 차수별/월별 그룹화
        # 차수 식별자: schedule_raw
        rounds_dict = {}
        for r in filtered:
            sched = r["schedule_raw"]
            if sched not in rounds_dict:
                rounds_dict[sched] = {
                    "schedule_raw": sched,
                    "start_date": r["start_date"],
                    "end_date": r["end_date"],
                    "month": r["month"],
                    "round_name": r["round_name"],
                    "display_period": r["display_period"],
                    "respondent_count": r["respondent_count"],
                    "course_satisfaction": r["course_satisfaction"],
                    "nps": r["nps"],
                    "records": []
                }
            rounds_dict[sched]["records"].append(r)
            rounds_dict[sched]["respondent_count"] = max(rounds_dict[sched]["respondent_count"], r["respondent_count"])

            # 5점 만점 정상 만족도 우선 채택
            sat = r["course_satisfaction"]
            if 1.0 <= sat <= 5.0 or rounds_dict[sched]["course_satisfaction"] == 0:
                rounds_dict[sched]["course_satisfaction"] = sat

            # NPS 정상값 채택
            n_val = r["nps"]
            if n_val > 10.0 or n_val < 0 or rounds_dict[sched]["nps"] == 0:
                rounds_dict[sched]["nps"] = n_val

        ordered_rounds = sorted(list(rounds_dict.values()), key=lambda x: x["start_date"])

        # 4. 차수별 누적 평균 계산
        cumulative_scores = {
            "teaching_expertise": [],
            "delivery_skill": [],
            "practical_use": [],
            "textbook_quality": [],
            "course_satisfaction": [],
            "nps": []
        }

        # 각 시트(차수)별 상세 데이터 구성
        sheet_data_list = []
        
        # 차수별 누적 추이용 히스토리
        trend_history = {
            "labels": [],
            "satisfaction": [],
            "nps": []
        }

        for idx, round_item in enumerate(ordered_rounds):
            trend_history["labels"].append(round_item["display_period"] or f"{round_item['month']}월")
            trend_history["satisfaction"].append(round_item["course_satisfaction"])
            trend_history["nps"].append(round_item["nps"])

            # 강사별 이번 차수 및 누적 차수 계산
            instructor_rows = []
            for rec in round_item["records"]:
                inst = rec["instructor_name"]
                # 해당 강사의 과거 모든 차수(이번 차수 포함) 점수 수집
                past_inst_records = []
                for p_round in ordered_rounds[:idx+1]:
                    for p_rec in p_round["records"]:
                        if p_rec["instructor_name"] == inst:
                            past_inst_records.append(p_rec)

                # 누적 평균
                count = len(past_inst_records)
                cum_expertise = round(sum(p["teaching_expertise"] for p in past_inst_records) / count, 3) if count > 0 else rec["teaching_expertise"]
                cum_delivery = round(sum(p["delivery_skill"] for p in past_inst_records) / count, 3) if count > 0 else rec["delivery_skill"]
                cum_practical = round(sum(p["practical_use"] for p in past_inst_records) / count, 3) if count > 0 else rec["practical_use"]
                cum_textbook = round(sum(p["textbook_quality"] for p in past_inst_records) / count, 3) if count > 0 else rec["textbook_quality"]

                # 커리큘럼 텍스트
                curr_key = f"{course_name}_{inst}"
                curriculum_text = self.curriculums.get(curr_key, f"과정: {course_name}\n강의시간: {rec['hours']}시간")

                instructor_rows.append({
                    "instructor_name": inst,
                    "curriculum": curriculum_text,
                    "hours": rec["hours"],
                    "current": {
                        "teaching_expertise": rec["teaching_expertise"], # 강의내용
                        "delivery_skill": rec["delivery_skill"],         # 전달능력
                        "practical_use": rec["practical_use"],           # 실무활용도
                        "textbook_quality": rec["textbook_quality"]       # 교재완성도
                    },
                    "cumulative": {
                        "teaching_expertise": cum_expertise,
                        "delivery_skill": cum_delivery,
                        "practical_use": cum_practical,
                        "textbook_quality": cum_textbook
                    }
                })

            # 차트 데이터 (차트 0~5)
            # 차트 0: 교육과정내용 만족도 추이 (현재 차수까지의 추이)
            chart0_data = {
                "categories": list(trend_history["labels"]),
                "values": list(trend_history["satisfaction"])
            }
            # 차트 1: 추천지수(NPS) 추이 (현재 차수까지의 추이)
            chart1_data = {
                "categories": list(trend_history["labels"]),
                "values": list(trend_history["nps"])
            }
            
            # 차트 2~5: 해당 차수의 공통 설문 데이터 (동일 차수 복수 강사 시 중복 합산 방지 -> 대표 레코드 사용)
            primary_rec = round_item["records"][0] if round_item["records"] else {}

            # 불편요소
            comp_keys = ["정보", "절차", "운영", "환경", "일정", "기타"]
            comp_vals = [primary_rec.get("complaints", {}).get(k, 0) for k in comp_keys]
            
            # 희망 형태
            form_keys = ["오프라인", "비대면", "이러닝", "플립러닝"]
            form_vals = [primary_rec.get("preferred_formats", {}).get(k, 0) for k in form_keys]

            # 직급
            pos_keys = ["사원", "대리", "과장", "차장", "부팀장", "임원"]
            pos_vals = [primary_rec.get("positions", {}).get(k, 0) for k in pos_keys]

            # 수강동기 (과정정보 출처)
            mot_keys = ["SNS", "이메일", "홈페이지", "인쇄물", "인터넷", "담당자추천", "동료추천", "기타"]
            mot_vals = [primary_rec.get("motives", {}).get(k, 0) for k in mot_keys]

            # 주관식 코멘트 가져오기
            comment_key = f"{course_name}_{round_item['schedule_raw']}_{instructor_name or 'all'}"
            round_comments = self.comments.get(comment_key, {
                "instructor_feedback": [],
                "content_feedback": [],
                "recommend_feedback": [],
                "instructor_sentiment": [],
                "content_sentiment": [],
                "recommend_sentiment": []
            })

            # 시트명 결정 (예: 3월, 7월 등)
            sheet_title = f"{round_item['month']}월" if round_item["month"] else f"{idx+1}차"
            # 중복 방지
            existing_sheet_titles = [s["sheet_title"] for s in sheet_data_list]
            if sheet_title in existing_sheet_titles:
                sheet_title = f"{sheet_title}({round_item['round_name'] or idx+1})"

            sheet_data_list.append({
                "sheet_title": sheet_title,
                "schedule_raw": round_item["schedule_raw"],
                "display_period": round_item["display_period"],
                "respondent_count": round_item["respondent_count"],
                "course_name": course_name,
                "instructors": instructor_rows,
                "charts": {
                    "chart0_satisfaction_trend": chart0_data,
                    "chart1_nps_trend": chart1_data,
                    "chart2_complaints": {"categories": comp_keys, "values": comp_vals},
                    "chart3_preferred_format": {"categories": form_keys, "values": form_vals},
                    "chart4_positions": {"categories": pos_keys, "values": pos_vals},
                    "chart5_motives": {"categories": mot_keys, "values": mot_vals},
                },
                "comments": round_comments,
                "comment_key": comment_key
            })

        return {
            "course_name": course_name,
            "instructor_name": instructor_name,
            "mode": mode,
            "sheets": sheet_data_list
        }

    def save_comments(self, comment_key: str, comments: Dict[str, Any]):
        self.comments[comment_key] = comments
        self.save_comments_store()

    def save_curriculum(self, course_name: str, instructor_name: str, curriculum_text: str):
        curr_key = f"{course_name}_{instructor_name}"
        self.curriculums[curr_key] = curriculum_text
        self.save_curriculum_store()


# 싱글톤 인스턴스
data_manager = DataManager()
